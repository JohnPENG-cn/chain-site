"""HBM 链唯一构建脚本（v0.4 起）。重跑不会丢内容：
  手写内容（来源、价格点、工艺表、玩家、证伪条件、更新记录……）只存在 hbm/content.json，人工编辑它；
  数字（需求、季度产能、供给、历史供给、回测、毛利）全部由本脚本从带来源的锚点算出。
输出：data/hbm.json、data/hbm_x.json、hbm/hbm_model.xlsx。
依赖：cowos_data/model_out.json（需求与 CoWoS 共用同一套 KGI 出货表），hbm_data/backtest.json，hbm_data/raw/mu_gm.json。
运行：python3 build.py（仓库根目录的 build.sh 会按 cowos/model.py → hbm/build.py → cowos/build.py 的顺序跑）。"""
import json,os,re,sys
from datetime import date,timedelta
R=os.path.dirname(os.path.dirname(os.path.abspath(__file__)));P=lambda *a:os.path.join(R,*a)
C=json.load(open(P('hbm','content.json')))
M={k:v for k,v in C.items() if k!='x'};X=C['x']
M['updated']='2026-10-15'
# ---------- 需求：与 CoWoS 页共用 KGI 出货表（出货 × 每封装 HBM 容量），源55/56
VN={'英伟达':'英伟达','谷歌TPU':'谷歌TPU','AWS':'AWS','AMD':'AMD'}
cw=json.load(open(P('cowos_data','model_out.json')))
M['demand']=[{'y':r['y'],'v':VN.get(r['v'],'其他/中国'),'chip':f"{r['v']} {r['chip']}",'units':r['units'],'gb':r['gb'],'src':r['src']+'（经 CoWoS 页）'} for r in cw['demand']]
# ---------- 季度产能（千片/月，季末），锚点见 anchors
Q=[f'{q}Q{y%100}' for y in range(2023,2028) for q in range(1,5)]
def interp(a,b,n=4): return [round(a+(b-a)*(i+1)/n,1) for i in range(n)]
ss={'SK hynix':[15,20,35,45,60,80,110,135],'Samsung':[15,20,35,45,45,45,80,115],'Micron':[4,4,4,4,10,15,25,30],'CXMT':[0]*8}
ye={'SK hynix':[175,230,270],'Samsung':[150,170,260],'Micron':[65,100,160],'CXMT':[0,10,25]}   # 2025/26/27 年末，UBS（源63,64）；长鑫为估
cap={'q':Q,'kind':['图读数']*8+['插值','插值','插值','锚点']*3}
for k in ss:
    v=ss[k][:];prev=v[-1]
    for t in ye[k]: v+=interp(prev,t);prev=t
    cap[k]=v
cap['anchors']=[
 {'p':'1Q23–4Q24','t':'三星证券 2024-03 报告 HBM Capa 季度图（图上读数，误差约 ±5K；2024 为当时预测）','src':24},
 {'p':'2024 年底','t':'TrendForce 2024-03：三星约 130K、海力士约 120K（含 TSV），全行业约 250K/月；美光 4Q24 约 25K（工商时报）','src':'25,26'},
 {'p':'2025–2027 年底','t':'UBS（2026-10 汇总）：海力士 175 → 230 → 270K，三星 150 → 170 → 260K，美光 65 → 100 → 160K，三家合计 390 → 500 → 690K/月；海力士 2026 年底 23 万、2027 年底 27 万另见首尔经济转述','src':'63,64'},
 {'p':'交叉验证','t':'三星：2026 年 HBM 平均月投片约 18 万、2027 年约 25 万（首尔经济，经 TrendForce），比 UBS 口径高约 10%；ETNews 称美光 2025 年实际约 40–50K、2026 年底约 100K，2025 年比 UBS 低。三星、海力士 2026 年各 15–20 万（ETNews）','src':'66,7,28'},
 {'p':'长鑫','t':'2026 年底 10K、2027 年底 25K 有效产能为估（良率约 25–30%）','src':'22'}]
M['cap']=cap
def qavg(k,y): i=Q.index(f'1Q{y%100}');return sum(cap[k][i:i+4])/4
MK=['SK hynix','Samsung','Micron','CXMT']
# ---------- 供给：年均月投片（= 季末产能 4 季平均，与约束页口径一致）× 12 × 每片有效 GB × 稼动率
M['supply']={str(y):{k:round(qavg(k,y),1) for k in MK} for y in (2025,2026,2027)}
M['gbw']={'2025':750,'2026':880,'2027':880}   # 校准：iM 证券 2026-06（源65）2025 年产量 26.3 亿 GB、2026 年 44.4 亿 GB
M['ut']={'2025':.85,'2026':.9,'2027':.9}
def sup(y): return sum(M['supply'][str(y)].values())*12*M['gbw'][str(y)]*M['ut'][str(y)]/1e6
M['calib']={'note':'每片有效 GB 不是原厂披露值，而是用券商汇总的产量反推：2025 年 iM 证券称业界产量 26.3 亿 GB、2026 年预计 44.4 亿 GB [源65]；用 UBS 年末产能 [源63]插值得到的年均投片去除，得到 2025 年约 750、2026 年约 880 GB/片（稼动率 85%/90%，估）。2027 年 HBM4/4E die 更大、16 层良率更低 [源1]，先按与 2026 年持平处理。旧版 600/650/600 是\"让 2026 年供需接近平衡\"倒推的，加上产能锚点偏低，2026 年供给少算了约三分之一（2.9 对 4.4 EB）。',
 'check':[['2025 产量','%.2f'%sup(2025),'2.63','iM 证券 2026-06（源65）'],['2026 产量','%.2f'%sup(2026),'4.44','iM 证券 2026-06（源65）'],['2026 需求','%.2f'%(sum(r['units']*r['gb'] for r in M['demand'] if r['y']==2026)/1000),'4.23（iM）/ >3.75（TrendForce 2025-05）','iM 证券（源65）/ TrendForce（源37）'],['2027 供给增速','%+.0f%%'%((sup(2027)/sup(2026)-1)*100),'+45–50%','KGI（源55）']]}
# ---------- 2023/24 历史供给（估）
hs={}
for y,g,u in[(2023,350,.85),(2024,450,.80)]:
    w=sum(qavg(k,y) for k in MK)
    hs[str(y)]={'wafers':round(w,1),'gbw':g,'ut':u,'s':round(w*12*g*u/1e6,3)}
M['hist_supply']=hs
M['hist_note']='2023/24 供给 = 季度产能均值 × 12 × 每片有效 GB（350/450，估）× 稼动率（85%%/80%%，估）。交叉验证：2024/2023 = %.1f 倍，TrendForce 2024-03 称 2024 年 HBM 位元供给增长约 260%%（3.6 倍）。'%(hs['2024']['s']/hs['2023']['s'])
# ---------- 回测、毛利
M['bt']=json.load(open(P('hbm_data','backtest.json')))
def cq(s):
    d=date.fromisoformat(s);d=d-timedelta(days=10) if d.day<10 else d;return f'{d.year}Q{(d.month-1)//3+1}'
M['mu_gm']={cq(k):v for k,v in json.load(open(P('hbm_data','raw','mu_gm.json'))).items()}
# ---------- 校验：引用的来源都存在、来源都有链接
ids={s['id'] for s in M['sources']}
assert len(ids)==len(M['sources']),'来源编号重复'
assert all(s.get('u','').startswith('http') for s in M['sources']),'来源缺链接'
txt=json.dumps(C,ensure_ascii=False)+open(P('hbm.js')).read()+open(P('hbm.html')).read()
miss=sorted({int(n) for g in re.findall(r'源([\d,，]+)',txt) for n in re.split('[,，]',g) if n} - ids)
assert not miss,f'引用了不存在的来源：{miss}'
M=dict(sorted(M.items()))
json.dump(M,open(P('data','hbm.json'),'w'),ensure_ascii=False,separators=(',',':'))
json.dump(X,open(P('data','hbm_x.json'),'w'),ensure_ascii=False,indent=0)
# ---------- Excel（黄色格可改；公式与网页一致）
try:
    import openpyxl;from openpyxl.styles import Font,PatternFill
    wb=openpyxl.Workbook();b=Font(bold=True);inp=PatternFill('solid',fgColor='FFF2CC');ws=wb.active;ws.title='说明'
    for r in [f'HBM 供需模型 v0.4（{M["updated"]}，由 hbm/build.py 生成，勿手改）','黄色格=可改假设；其余为公式。1 EB = 10 亿 GB。',
      '需求 = 出货（百万颗）× 每封装 HBM（GB），与 CoWoS 页共用 KGI 出货表','供给 = 年均月投片（千片）× 12 × 每片有效 GB × 稼动率',M['calib']['note']]: ws.append([r])
    wd=wb.create_sheet('需求');wd.append(['年份','厂商','芯片','出货(百万)','每封装HBM(GB)','需求(PB)','来源']);[setattr(c,'font',b) for c in wd[1]]
    for i,r in enumerate(M['demand'],start=2):
        wd.append([r['y'],r['v'],r['chip'],r['units'],r['gb'],f'=D{i}*E{i}',r['src']]);wd[f'D{i}'].fill=inp;wd[f'E{i}'].fill=inp
    n=len(M['demand'])+1
    wsu=wb.create_sheet('供给');wsu.append(['年份','厂商','年均月投片(千片)','每片有效GB','稼动率','年供给(PB)']);[setattr(c,'font',b) for c in wsu[1]];r=2
    for y,m in M['supply'].items():
        for k,v in m.items():
            wsu.append([int(y),k,v,M['gbw'][y],M['ut'][y],f'=C{r}*12*D{r}*E{r}/1000']);[setattr(wsu[f'{c}{r}'],'fill',inp) for c in 'CDE'];r+=1
    wq=wb.create_sheet('季度产能');wq.append(['季度','类型']+MK);[setattr(c,'font',b) for c in wq[1]]
    for i,q in enumerate(Q): wq.append([q,cap['kind'][i]]+[cap[k][i] for k in MK])
    wbal=wb.create_sheet('平衡表');wbal.append(['年份','需求(EB)','供给(EB)','缺口/需求']);[setattr(c,'font',b) for c in wbal[1]]
    for i,y in enumerate([2025,2026,2027],start=2):
        wbal.append([y,f'=SUMIF(需求!A2:A{n},A{i},需求!F2:F{n})/1000',f'=SUMIF(供给!A2:A{r-1},A{i},供给!F2:F{r-1})/1000',f'=(C{i}-B{i})/B{i}'])
    wc=wb.create_sheet('校准');wc.append(['项目','本模型','外部口径','来源']);[wc.append(x) for x in M['calib']['check']]
    wsrc=wb.create_sheet('来源');wsrc.append(['编号','说明','链接']);[wsrc.append([s['id'],s['t'],s['u']]) for s in M['sources']]
    wb.save(P('hbm','hbm_model.xlsx'))
except ImportError: print('openpyxl 未安装，跳过 Excel')
for y in (2025,2026,2027):
    d=sum(r['units']*r['gb'] for r in M['demand'] if r['y']==y)/1000;s=sup(y);print(y,'需求 %.2f 供给 %.2f 缺口 %+.0f%%'%(d,s,(s-d)/d*100))
print('ok · 来源',len(ids))
