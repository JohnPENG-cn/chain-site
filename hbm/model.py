import openpyxl
from openpyxl.styles import Font, PatternFill
import matplotlib; matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager as fm
fm.fontManager.addfont('/tmp/wenkai.ttf'); FN=fm.FontProperties(fname='/tmp/wenkai.ttf').get_name()
plt.rcParams['font.family']=FN; plt.rcParams['axes.unicode_minus']=False
# demand: (year, chip, units M, GB/chip, source)
D=[
(2023,'H100等Hopper',1.5,80,'估算'),(2023,'其他(TPU v5/MI300等)',0.8,64,'估算'),
(2024,'H100/H200',3.0,100,'估算'),(2024,'TPU v5/v6',1.5,95,'估算'),(2024,'AMD MI300',0.4,192,'估算'),(2024,'其他ASIC',0.6,64,'估算'),
(2025,'Blackwell B200/GB200',4.0,192,'估算'),(2025,'H200/H20',1.5,141,'估算'),(2025,'Google TPU',2.5,192,'估算'),(2025,'AWS Trainium2',1.5,96,'估算'),(2025,'AMD MI325/355',0.5,288,'估算'),
(2026,'Blackwell B200/B300(GB300)',5.45,250,'出货:Hilo转述券商~5.4-5.5M;容量:B200 192/B300 288按4:6估算'),
(2026,'Rubin',1.5,288,'出货:Magica报道下调至~1.5M;容量TrendForce称与上代相近'),
(2026,'H200等Hopper',0.5,141,'估算'),
(2026,'Google TPU',3.7,192,'出货:JPMorgan 3.7M;容量:Ironwood 192GB'),
(2026,'AWS Trainium2/3',1.5,120,'估算(Trn3 144GB)'),
(2026,'AMD MI355/MI450',0.5,288,'估算'),(2026,'其他ASIC(Meta/MS)',0.5,128,'估算'),(2026,'中国AI芯片',0.7,64,'估算'),
(2027,'Rubin',5.5,288,'出货:券商Rubin+Ultra近7M,拆分为估算'),
(2027,'Rubin Ultra',1.5,384,'容量:TrendForce 384GB/GPU'),
(2027,'Blackwell尾量',1.5,288,'估算'),
(2027,'Google TPU',5.0,250,'出货:JPMorgan 5M;容量估算'),
(2027,'AWS Trainium3',2.5,144,'估算'),
(2027,'AMD MI450',1.0,432,'估算'),(2027,'其他ASIC',1.0,192,'估算'),(2027,'中国AI芯片',1.0,96,'估算'),
]
# supply: year -> {maker: avg wafers/month (k)}
S={2025:{'SK hynix':150,'Samsung':110,'Micron':40,'CXMT':0},
   2026:{'SK hynix':170,'Samsung':165,'Micron':70,'CXMT':10},
   2027:{'SK hynix':230,'Samsung':250,'Micron':110,'CXMT':25}}
GBW={2025:600,2026:650,2027:600}; UT={2025:0.85,2026:0.9,2027:0.9}
wb=openpyxl.Workbook(); b=Font(bold=True); inp=PatternFill('solid',fgColor='FFF2CC')
ws=wb.active; ws.title='说明'
for r in ["HBM 供需模型 v0.1（2026-10-14，Kiko）","黄色格=可改假设；其余为公式。单位：1 EB = 10亿 GB。",
 "需求=出货(百万颗)×每颗HBM容量(GB)；供给=平均月投片(千片)×12×每片有效GB×稼动率",
 "每片有效GB为自估：以TrendForce'HBM占DRAM投片22%但位元仅9%'及美光'HBM耗片约DDR5的3倍'反推，并校准到2026年供需接近平衡（与TrendForce 2026'价格暂停'口径一致）",
 "2023-2024 供给缺少可靠公开数据，未建模。所有出货/容量未标来源者均为自估。"]: ws.append([r])
wd=wb.create_sheet('需求'); wd.append(['年份','芯片','出货(百万)','每颗HBM(GB)','需求(PB)','来源/假设']);[setattr(c,'font',b) for c in wd[1]]
for i,(y,c,u,g,s) in enumerate(D,start=2):
    wd.append([y,c,u,g,f'=C{i}*D{i}',s]); wd[f'C{i}'].fill=inp; wd[f'D{i}'].fill=inp
n=len(D)+1
wsu=wb.create_sheet('供给'); wsu.append(['年份','厂商','平均月投片(千片)','每片有效GB','稼动率','年供给(PB)']);[setattr(c,'font',b) for c in wsu[1]]
r=2
for y,m in S.items():
    for k,v in m.items():
        wsu.append([y,k,v,GBW[y],UT[y],f'=C{r}*12*D{r}*E{r}/1000']); [setattr(wsu[f'{c}{r}'],'fill',inp) for c in 'CDE']; r+=1
wbal=wb.create_sheet('平衡表'); wbal.append(['年份','需求(EB)','供给(EB)','缺口(供-需,EB)','缺口/需求','判断']);[setattr(c,'font',b) for c in wbal[1]]
res={}
for i,y in enumerate([2023,2024,2025,2026,2027],start=2):
    dem=sum(u*g for yy,_,u,g,_ in D if yy==y)/1000
    sup=sum(v*12*GBW[y]*UT[y]/1000 for v in S[y].values())/1000 if y in S else None
    res[y]=(dem,sup)
    if sup: wbal.append([y,f"=SUMIF(需求!A2:A{n},A{i},需求!E2:E{n})/1000",f"=SUMIF(供给!A2:A{r-1},A{i},供给!F2:F{r-1})/1000",f'=C{i}-B{i}',f'=D{i}/B{i}',''])
    else: wbal.append([y,f"=SUMIF(需求!A2:A{n},A{i},需求!E2:E{n})/1000",'未建模','','','公开口径:全年售罄/紧缺'])
wb.create_sheet('来源').append(['见 hbm_chain.md 附录 来源清单'])
wb.save('hbm_model.xlsx')
for y,(d,s) in res.items(): print(y,round(d,2),s and round(s,2), s and f"{(s-d)/d:+.0%}")
# chart1 balance
ys=[2025,2026,2027]; fig,ax=plt.subplots(figsize=(8,5),dpi=150)
ax.bar([y-0.2 for y in ys],[res[y][0] for y in ys],0.4,label='需求',color='#d9534f')
ax.bar([y+0.2 for y in ys],[res[y][1] for y in ys],0.4,label='供给',color='#5b8def')
for y in ys:
    d,s=res[y]; ax.text(y,max(d,s)+0.1,f"{(s-d)/d:+.0%}",ha='center',fontsize=14)
ax.set_xticks(ys); ax.set_ylabel('EB（10亿GB）'); ax.set_title('HBM 供需平衡（Kiko自估基准情形）',fontsize=15); ax.legend(loc='upper left')
fig.text(0.01,0.01,'来源：出货参考JPMorgan/券商转述，产能参考ETNews/Chosun/TrendForce；每片有效GB等为自估',fontsize=8)
plt.tight_layout(); plt.savefig('chart_balance.png')
# chart2 demand mix 2027 vs 2026
fig,ax=plt.subplots(figsize=(8,5),dpi=150)
import collections
for j,y in enumerate([2026,2027]):
    agg=collections.OrderedDict()
    for yy,c,u,g,_ in D:
        if yy!=y: continue
        k='英伟达' if any(x in c for x in['Blackwell','Rubin','H100','H200','Hopper']) else ('谷歌TPU' if 'TPU' in c else ('AWS' if 'Trainium' in c else ('AMD' if 'AMD' in c else '其他/中国')))
        agg[k]=agg.get(k,0)+u*g/1000
    left=0
    cols={'英伟达':'#76b900','谷歌TPU':'#4285f4','AWS':'#ff9900','AMD':'#ed1c24','其他/中国':'#999'}
    for k,v in agg.items():
        ax.barh(j,v,left=left,color=cols[k],label=k if j==0 else None); 
        if v>0.25: ax.text(left+v/2,j,f'{k}\n{v:.1f}',ha='center',va='center',color='white',fontsize=10)
        left+=v
ax.set_yticks([0,1]); ax.set_yticklabels(['2026','2027']); ax.set_xlabel('EB'); ax.set_title('HBM 需求由谁拉动（自估）',fontsize=15)
plt.tight_layout(); plt.savefig('chart_demand_mix.png')
# chart3 margins
fig,ax=plt.subplots(figsize=(8,5),dpi=150)
lab=['SK海力士\n营业利润率 2Q25','SK海力士\n营业利润率 2Q26','美光毛利率\nFQ4 25','美光毛利率\nFQ4 26','台积电\n营业利润率 2Q26']
val=[41,76,45.7,87.0,60.3]; ax.bar(lab,val,color=['#bbb','#d9534f','#bbb','#d9534f','#5b8def'])
for i,v in enumerate(val): ax.text(i,v+1,f'{v}%',ha='center',fontsize=12)
ax.set_ylim(0,100); ax.set_title('利润堆在存储：内存厂利润率已超过台积电',fontsize=15)
fig.text(0.01,0.01,'来源：SK海力士2Q26财报、美光FQ4 2026财报(Non-GAAP)、Yonhap(台积电)',fontsize=8)
plt.tight_layout(); plt.savefig('chart_margin.png')
